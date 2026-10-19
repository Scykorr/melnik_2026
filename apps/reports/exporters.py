from __future__ import annotations

from io import BytesIO
from pathlib import Path

from django.conf import settings
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

FONT_PATH = Path(settings.BASE_DIR) / "static" / "fonts" / "DejaVuSans.ttf"
_FONT_READY = False


def _ensure_font() -> str:
    global _FONT_READY
    if FONT_PATH.exists():
        if not _FONT_READY:
            pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_PATH)))
            _FONT_READY = True
        return "DejaVu"
    return "Helvetica"


def _lines(report) -> list[str]:
    body = report.body_json or {}
    auto = body.get("auto_text") or {}
    approved = body.get("approved_text") or {}
    lines = [
        settings.PROJECT_TITLE,
        report.title,
        f"Территория: {report.region.name}",
        f"Период: {report.period_start.isoformat()} — {report.period_end.isoformat()}",
        f"Версия: {report.version}. Статус: {report.get_status_display()}",
        f"Методика: {report.method_version}. Шаблон: {report.template_version}",
        f"Контрольная сумма входов: {report.input_hash}",
        "",
        "Сформировано автоматически",
        auto.get("summary", ""),
        auto.get("topics", ""),
        auto.get("dynamics", ""),
        auto.get("limitations", ""),
        "",
        "Текст, утверждённый аналитиком",
        approved.get("summary", "") or "Раздел пока не отредактирован аналитиком",
        approved.get("topics", ""),
        approved.get("dynamics", ""),
        approved.get("facts", ""),
        approved.get("limitations", ""),
        "",
        "Подтверждённые факты",
    ]
    for fact in body.get("facts") or []:
        lines.append(f"- {fact.get('title')} ({fact.get('topic')})")
        for source in fact.get("sources") or []:
            lines.append(f"  источник: {source.get('name')} {source.get('url')}")
    lines.append("")
    lines.append("Трудозатраты T0, T1 и K_T: не измерено.")
    lines.append("Материал не содержит персональных профилей и не оценивает взгляды отдельных людей.")
    return [line for line in lines if line is not None]


def render_docx(report) -> bytes:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for index, line in enumerate(_lines(report)):
        if index < 2:
            document.add_heading(line, level=index + 1)
        else:
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_pdf(report) -> bytes:
    font = _ensure_font()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title=report.title)
    style = ParagraphStyle(name="Body", fontName=font, fontSize=11, leading=15)
    story = []
    for line in _lines(report):
        safe = (
            line.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        story.append(Paragraph(safe or " ", style))
        story.append(Spacer(1, 6))
    doc.build(story)
    return buffer.getvalue()
